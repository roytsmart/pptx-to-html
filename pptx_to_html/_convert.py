"""Turn a PowerPoint deck into a single web page."""

import dataclasses
import fractions
import html
import importlib.resources
import io
import pathlib
import posixpath
import subprocess
import tempfile
import warnings
from collections.abc import Callable, Iterator
from typing import TypeAlias

import PIL.Image
import pptx
from lxml import etree
from pptx.opc.package import Part
from pptx.presentation import Presentation
from pptx.slide import Slide, SlideLayout, SlideMaster

from ._fonts import font_faces, font_stack
from ._math import mathml
from ._text import Inheritance, TextRenderer
from ._theme import Color, Theme, fill_color, line_style
from ._xml import NS, find, findall, integer, local, qn

__all__ = [
    "ConversionWarning",
    "Reencode",
    "convert",
]


class ConversionWarning(UserWarning):
    """Something on a slide that the page cannot show the way PowerPoint does."""


@dataclasses.dataclass
class Reencode:
    """How to shrink the movies on a deck for the web."""

    ffmpeg: str = "ffmpeg"
    """The ``ffmpeg`` executable to use."""

    crf: int = 26
    """The H.264 constant rate factor: higher is smaller and blurrier."""

    max_width: int = 1920
    """Movies wider than this many pixels are scaled down to it."""


@dataclasses.dataclass
class _Box:
    """Where a shape sits on its slide, in EMU."""

    x: float
    y: float
    width: float
    height: float
    rotation: float = 0
    flip_h: bool = False
    flip_v: bool = False


_Transform: TypeAlias = Callable[[_Box], _Box]
"""Moves a shape's box from the group it is in onto the slide."""


def convert(
    deck: "str | pathlib.Path",
    output: "str | pathlib.Path",
    title: "str | None" = None,
    subtitle: "str | None" = None,
    media: "dict[str, str] | None" = None,
    reencode: "Reencode | None" = None,
    include_hidden: bool = False,
    notes: bool = False,
    relink: "dict[str, str] | None" = None,
) -> pathlib.Path:
    """
    Write a deck out as a web page: ``index.html`` and a ``media`` folder.

    Parameters
    ----------
    deck
        The ``.pptx`` file to convert.
    output
        The folder to write the page into. It is created if need be.
    title
        The heading of the page. By default, the deck's title, or the text
        of the title on its first slide.
    subtitle
        A line under the heading, if any.
    media
        Addresses to use instead of the copies of a picture or movie
        embedded in the deck, by the name the deck gives the file (like
        ``media1.mp4``). Useful for pointing at a sharper original.
    reencode
        If given, movies are re-encoded with ``ffmpeg`` to make them
        smaller.
    include_hidden
        Whether to include slides hidden in the slide show.
    notes
        Whether to show each slide's speaker notes under it. Off by default,
        since notes are often not meant for an audience.
    relink
        Folders that files the deck links to have moved from, and where they
        are now, like ``{"C:/Users/old": "C:/Users/new"}``. A linked file
        found on this computer is copied next to the page like an embedded
        one.

    Returns
    -------
    The path of the ``index.html`` written.
    """
    output = pathlib.Path(output)
    output.mkdir(parents=True, exist_ok=True)
    presentation = pptx.Presentation(str(deck))
    width = int(presentation.slide_width or 12192000)
    height = int(presentation.slide_height or 6858000)

    library = _Media(output / "media", media or {}, reencode, relink or {})
    slides = [
        s for s in presentation.slides if include_hidden or s.element.get("show") != "0"
    ]
    master = presentation.slide_masters[0]
    theme = Theme(master)
    default_color = theme.color(_scheme("tx1"))

    sections = []
    any_notes = False
    for number, slide in enumerate(slides, start=1):
        converter = _SlideConverter(
            presentation=presentation,
            slide=slide,
            width=width,
            height=height,
            library=library,
            number=number,
            default_color=default_color,
        )
        remarks = _notes(slide) if notes else ""
        any_notes = any_notes or bool(remarks)
        sections.append(
            f'<section class="slide" id="slide-{number}">'
            f'<div class="stage"{converter.background()}>{"".join(converter.pieces())}</div>'
            f'<div class="number">{number} / {len(slides)}</div>{remarks}</section>'
        )

    if title is None:
        title = presentation.core_properties.title or _first_title(slides) or "Slides"

    ratio = fractions.Fraction(width, height).limit_denominator(100)
    template = importlib.resources.files(__package__).joinpath("template.html").read_text("utf-8")
    page = (
        template.replace("{{title}}", html.escape(title))
        .replace(
            "{{subtitle}}",
            f"\n  <p>{html.escape(subtitle)}</p>" if subtitle else "",
        )
        .replace("{{aspect}}", f"{ratio.numerator} / {ratio.denominator}")
        .replace("{{font}}", font_stack(theme.font_minor))
        .replace("{{color}}", default_color.css() if default_color else "#000000")
        .replace("{{controls}}", '\n  <button id="notes">Hide notes</button>' if any_notes else "")
        .replace("{{slides}}", "\n".join(sections))
    )
    page = page.replace("{{fonts}}", font_faces(page))
    index = output / "index.html"
    index.write_text(page, encoding="utf-8", newline="\n")
    return index


class _SlideConverter:
    """Turns the shapes on one slide into positioned elements."""

    def __init__(
        self,
        presentation: Presentation,
        slide: Slide,
        width: int,
        height: int,
        library: "_Media",
        number: int,
        default_color: "Color | None",
    ):
        self.presentation = presentation
        self.slide = slide
        self.layout: SlideLayout = slide.slide_layout
        self.master: SlideMaster = self.layout.slide_master
        self.theme = Theme(self.master)
        self.width = width
        self.height = height
        self.library = library
        self.number = number
        self.z = 0
        self.default_color = default_color

    def warn(self, message: str) -> None:
        warnings.warn(f"slide {self.number}: {message}", ConversionWarning, stacklevel=4)

    def pieces(self) -> "Iterator[str]":
        """Every shape PowerPoint would draw on the slide, back to front."""
        layers: list[tuple[etree._Element, Part]] = []
        if self.slide.element.get("showMasterSp") != "0":
            if self.layout.element.get("showMasterSp") != "0":
                layers.append((self.master.element, self.master.part))
            layers.append((self.layout.element, self.layout.part))
        layers.append((self.slide.element, self.slide.part))

        for index, (owner, part) in enumerate(layers):
            tree = find(owner, "p:cSld/p:spTree")
            on_slide = index == len(layers) - 1
            for element, transform in _walk(tree, lambda box: box):
                # Placeholders on layouts and masters are only prompts for
                # the slides that use them, and are never drawn themselves.
                if not on_slide and _placeholder(element) is not None:
                    continue
                piece = self.piece(element, transform, part, on_slide)
                if piece:
                    yield piece
                    self.z += 1

    def piece(
        self,
        element: etree._Element,
        transform: "_Transform",
        part: Part,
        on_slide: bool,
    ) -> "str | None":
        kind = local(element)
        chain = self.placeholder_chain(element) if on_slide else []
        box = _box(element, chain)
        if box is None:
            self.warn(f"{_name(element)!r} has no position, so it was left out")
            return None
        box = transform(box)
        if kind == "pic":
            return self.picture(element, box, part)
        if kind == "sp":
            geometry = find(element, "p:spPr/a:prstGeom")
            if geometry is not None and geometry.get("prst") in ("line", "straightConnector1"):
                return self.line(element, box)
            return self.shape(element, box, chain, part)
        if kind == "cxnSp":
            return self.line(element, box)
        if kind == "graphicFrame":
            self.warn(f"{_name(element)!r} is a table, chart, or diagram, which are not converted")
        return None

    def placeholder_chain(self, element: etree._Element) -> "list[etree._Element]":
        """The layout and master placeholders a slide's placeholder is based on."""
        placeholder = _placeholder(element)
        if placeholder is None:
            return []
        kind = placeholder.get("type", "obj")
        index = placeholder.get("idx", "0")
        chain = []
        on_layout = _match(self.layout.element, kind, index)
        if on_layout is not None:
            chain.append(on_layout)
            layout_placeholder = _placeholder(on_layout)
            kind = layout_placeholder.get("type", "obj") if layout_placeholder is not None else kind
        master_kind = {"ctrTitle": "title", "subTitle": "body", "obj": "body"}.get(kind, kind)
        on_master = _match(self.master.element, master_kind, None)
        if on_master is not None:
            chain.append(on_master)
        return chain

    def style(self, box: _Box) -> str:
        styles = [
            f"z-index:{self.z}",
            f"left:{_fixed(100 * box.x / self.width)}%",
            f"top:{_fixed(100 * box.y / self.height)}%",
            f"width:{_fixed(100 * box.width / self.width)}%",
            f"height:{_fixed(100 * box.height / self.height)}%",
        ]
        transforms = []
        if box.rotation:
            transforms.append(f"rotate({box.rotation:g}deg)")
        if box.flip_h:
            transforms.append("scaleX(-1)")
        if box.flip_v:
            transforms.append("scaleY(-1)")
        if transforms:
            styles.append(f"transform:{' '.join(transforms)}")
        return ";".join(styles)

    def outline(self, element: etree._Element, box: _Box) -> "list[str]":
        """The CSS for a shape's fill-independent decorations: its outline and outline shape."""
        styles = []
        properties = find(element, "p:spPr")
        line = line_style(self.theme, properties, find(element, "p:style"))
        if line is not None:
            width, color = line
            styles.append(f"border:{self.cqh(width):.4f}cqh solid {color.css()}")
            styles.append("box-sizing:border-box")
        styles += self.geometry(find(properties, "a:prstGeom"), box)
        return styles

    def geometry(self, geometry: "etree._Element | None", box: _Box) -> "list[str]":
        """The CSS that cuts a box down to one of PowerPoint's preset shapes."""
        if geometry is None:
            return []
        preset = geometry.get("prst", "rect")
        adjust = {
            g.get("name"): int(g.get("fmla", "val 0").split()[-1])
            for g in findall(geometry, "a:avLst/a:gd")
        }
        shortest = min(box.width, box.height)
        if preset == "rect":
            return []
        if preset == "ellipse":
            return ["border-radius:50%"]
        if preset in ("roundRect", "snipRoundRect"):
            radius = adjust.get("adj", 16667) / 100000 * shortest
            return [f"border-radius:{self.cqh(radius):.4f}cqh"]
        if preset == "octagon":
            corner = adjust.get("adj", 29289) / 100000 * shortest
            dx = 100 * corner / box.width
            dy = 100 * corner / box.height
            points = [
                (dx, 0), (100 - dx, 0), (100, dy), (100, 100 - dy),
                (100 - dx, 100), (dx, 100), (0, 100 - dy), (0, dy),
            ]
        elif preset == "triangle":
            apex = adjust.get("adj", 50000) / 1000
            points = [(apex, 0), (100, 100), (0, 100)]
        elif preset == "rtTriangle":
            points = [(0, 0), (100, 100), (0, 100)]
        elif preset == "diamond":
            points = [(50, 0), (100, 50), (50, 100), (0, 50)]
        else:
            self.warn(f"the {preset!r} shape is drawn as a rectangle")
            return []
        polygon = ", ".join(f"{_number(x)}% {_number(y)}%" for x, y in points)
        return [f"clip-path:polygon({polygon})"]

    def cqh(self, emu: float) -> float:
        """A length in EMU as a percentage of the slide's height."""
        return 100 * emu / self.height

    def picture(self, element: etree._Element, box: _Box, part: Part) -> "str | None":
        blip = find(element, "p:blipFill/a:blip")
        poster = self.image_source(blip, part, element)
        movie = self.movie_source(element, part)
        # A picture swapped for a movie, such as a heavy animated GIF for the
        # MP4 it was made from, is played like one.
        if movie is None and poster is not None and poster.lower().endswith(_MOVIES):
            movie, poster = poster, None
        if movie is None and poster is None:
            return None

        outline = self.outline(element, box)
        crop = find(element, "p:blipFill/a:srcRect")
        sides = [integer(crop, side) / 100000 for side in ("l", "t", "r", "b")]

        if movie is not None:
            poster_attribute = f' poster="{html.escape(poster)}"' if poster else ""
            tag = "video"
            attributes = (
                f'src="{html.escape(movie)}"{poster_attribute} '
                "controls loop muted playsinline preload=\"metadata\""
            )
        else:
            tag = "img"
            description = _description(element)
            attributes = f'src="{html.escape(poster or "")}" alt="{html.escape(description)}"'
        close = "></video>" if tag == "video" else ">"

        if any(sides):
            left, top, right, bottom = sides
            width = 100 / (1 - left - right)
            height = 100 / (1 - top - bottom)
            inner = (
                f"position:absolute;width:{_fixed(width)}%;height:{_fixed(height)}%;"
                f"left:{_fixed(-left * width)}%;top:{_fixed(-top * height)}%"
            )
            wrapper = ";".join([self.style(box)] + outline)
            return (
                f'<div class="piece crop" style="{wrapper}">'
                f'<{tag} class="piece fill" style="{inner}" {attributes}{close}</div>'
            )
        style = ";".join([self.style(box)] + outline)
        return f'<{tag} class="piece fill" style="{style}" {attributes}{close}'

    def image_source(
        self,
        blip: "etree._Element | None",
        part: Part,
        element: etree._Element,
    ) -> "str | None":
        if blip is None:
            return None
        # PowerPoint keeps a vector original of an SVG alongside the bitmap
        # it draws for older versions of itself.
        svg = find(blip, "a:extLst/a:ext/asvg:svgBlip")
        for candidate, attribute in ((svg, "r:embed"), (blip, "r:embed"), (blip, "r:link")):
            source = self.source(part, _attribute(candidate, attribute), element)
            if source is not None:
                return source
        return None

    def movie_source(self, element: etree._Element, part: Part) -> "str | None":
        properties = find(element, "p:nvPicPr/p:nvPr")
        if find(properties, "a:audioFile") is not None:
            self.warn(f"{_name(element)!r} is a sound, which is not converted")
            return None
        embedded = find(properties, "p:extLst/p:ext/p14:media")
        linked = find(properties, "a:videoFile")
        for candidate, attribute in ((embedded, "r:embed"), (linked, "r:link")):
            source = self.source(part, _attribute(candidate, attribute), element, movie=True)
            if source is not None:
                return source
        return None

    def source(
        self,
        part: Part,
        rid: "str | None",
        element: etree._Element,
        movie: bool = False,
    ) -> "str | None":
        """
        The address on the page of the file a relationship points to.

        A file embedded in the deck is copied next to the page. One the deck
        only links to is left out with a warning, unless it was given a
        substitute by name.
        """
        if not rid or rid not in part.rels:
            return None
        relationship = part.rels[rid]
        if not relationship.is_external:
            return self.library.add(relationship.target_part, movie=movie)
        name = posixpath.basename(relationship.target_ref.replace("\\", "/"))
        if name in self.library.substitutes:
            return self.library.substitutes[name]
        found = self.library.locate(relationship.target_ref)
        if found is not None:
            return self.library.add_file(found, movie=movie)
        self.warn(
            f"{_name(element)!r} links to {relationship.target_ref} outside the deck "
            f"instead of embedding it, so it was left out (say where it moved with "
            f"--relink OLD=NEW, or give it a substitute with --media {name}=ADDRESS)"
        )
        return None

    def shape(
        self,
        element: etree._Element,
        box: _Box,
        chain: "list[etree._Element]",
        part: Part,
    ) -> "str | None":
        properties = find(element, "p:spPr")
        style_element = find(element, "p:style")
        fill = fill_color(self.theme, properties, style_element)
        outline = self.outline(element, box)
        has_line = any(s.startswith("border:") for s in outline)

        body = find(element, "p:txBody")
        paragraphs: list[str] = []
        styles = [self.style(box)]
        if body is not None:
            inheritance = self.inheritance(element, chain)
            renderer = TextRenderer(
                theme=self.theme,
                slide_height=self.height,
                default_color=self.default_color,
                hyperlink=lambda rid: _hyperlink(part, rid),
                warn=lambda message: self.warn(f"{_name(element)!r}: {message}"),
            )
            paragraphs = renderer.paragraphs(body, inheritance)
            if paragraphs:
                anchor = inheritance.body("anchor") or "t"
                justify = {"t": "flex-start", "ctr": "center", "b": "flex-end"}.get(anchor, "flex-start")
                insets = [
                    _integer_or(inheritance.body(name), default)
                    for name, default in (("tIns", 45720), ("rIns", 91440), ("bIns", 45720), ("lIns", 91440))
                ]
                styles.append(f"justify-content:{justify}")
                styles.append("padding:" + " ".join(f"{self.cqh(i):.3f}cqh" for i in insets))
                if inheritance.body("wrap") == "none":
                    styles.append("white-space:nowrap;width:max-content")
                if (inheritance.body("vert") or "horz") != "horz":
                    self.warn(f"{_name(element)!r} has vertical text, which is set horizontally")

        picture = self.picture_fill(element, properties, part)
        if not paragraphs and fill is None and not has_line and not picture:
            return None
        if fill is not None:
            styles.append(f"background:{fill.css()}")
        styles += outline
        classes = ["piece"]
        if paragraphs:
            classes.append("text")
        if picture:
            classes.append("crop")
        return (
            f'<div class="{" ".join(classes)}" style="{html.escape(";".join(styles))}">'
            f'{picture}{"".join(paragraphs)}</div>'
        )

    def picture_fill(self, element: etree._Element, properties: "etree._Element | None", part: Part) -> str:
        """
        A picture a shape is filled with, stretched over it. This is also how
        PowerPoint draws the stand-in for content an older reader cannot,
        such as a text box with an equation in it.
        """
        fill = find(properties, "a:blipFill")
        if fill is None:
            return ""
        source = self.image_source(find(fill, "a:blip"), part, element)
        if source is None:
            return ""
        # How far the picture's edges sit inside the shape's, in thousandths
        # of a percent; negative when the picture spills over them.
        rect = find(fill, "a:stretch/a:fillRect")
        left, top, right, bottom = (integer(rect, side) / 1000 for side in ("l", "t", "r", "b"))
        style = (
            f"position:absolute;left:{_fixed(left)}%;top:{_fixed(top)}%;"
            f"width:{_fixed(100 - left - right)}%;height:{_fixed(100 - top - bottom)}%"
        )
        return f'<img class="fill" style="{style}" src="{html.escape(source)}" alt="">'

    def inheritance(self, element: etree._Element, chain: "list[etree._Element]") -> Inheritance:
        placeholder = _placeholder(element)
        if placeholder is None:
            key = "p:otherStyle"
        elif placeholder.get("type") in ("title", "ctrTitle"):
            key = "p:titleStyle"
        else:
            key = "p:bodyStyle"
        inherited = [find(e, "p:txBody/a:lstStyle") for e in chain]
        inherited.append(find(self.master.element, f"p:txStyles/{key}"))
        inherited.append(find(self.presentation.element, "p:defaultTextStyle"))
        bodies = [find(e, "p:txBody/a:bodyPr") for e in [element] + chain]
        font_reference = find(element, "p:style/a:fontRef")
        return Inheritance(
            own_styles=[e for e in [find(element, "p:txBody/a:lstStyle")] if e is not None],
            inherited_styles=[e for e in inherited if e is not None],
            bodies=[e for e in bodies if e is not None],
            font_color=self.theme.color(font_reference),
        )

    def line(self, element: etree._Element, box: _Box) -> "str | None":
        properties = find(element, "p:spPr")
        line = line_style(self.theme, properties, find(element, "p:style"))
        if line is None:
            return None
        width, color = line
        x1, y1 = box.x, box.y
        x2, y2 = box.x + box.width, box.y + box.height
        if box.flip_h:
            x1, x2 = x2, x1
        if box.flip_v:
            y1, y2 = y2, y1
        markers = ""
        ends = ""
        for end, attribute in (("headEnd", "marker-start"), ("tailEnd", "marker-end")):
            head = find(find(properties, "a:ln"), f"a:{end}")
            if head is not None and head.get("type", "none") != "none":
                marker = f"arrow-{self.number}-{self.z}-{end}"
                markers += (
                    f'<marker id="{marker}" viewBox="0 0 10 10" refX="5" refY="5" '
                    'markerWidth="4" markerHeight="4" orient="auto-start-reverse">'
                    f'<path d="M0,0 L10,5 L0,10 z" fill="{color.css()}"/></marker>'
                )
                ends += f' {attribute}="url(#{marker})"'
        defs = f"<defs>{markers}</defs>" if markers else ""
        # Drawn in the slide's own units over the whole slide, so that the
        # line keeps its width however the slide is scaled.
        return (
            f'<svg class="piece" style="z-index:{self.z};left:0;top:0;width:100%;height:100%;overflow:visible" '
            f'viewBox="0 0 {self.width} {self.height}">{defs}'
            f'<line x1="{round(x1)}" y1="{round(y1)}" x2="{round(x2)}" y2="{round(y2)}" stroke="{color.css()}" '
            f'stroke-width="{width}"{ends}/></svg>'
        )

    def background(self) -> str:
        """A style attribute for the slide's stage, if its background is not plain white."""
        for owner, part in (
            (self.slide.element, self.slide.part),
            (self.layout.element, self.layout.part),
            (self.master.element, self.master.part),
        ):
            background = find(owner, "p:cSld/p:bg")
            if background is None:
                continue
            properties = find(background, "p:bgPr")
            reference = find(background, "p:bgRef")
            if properties is not None:
                picture = find(properties, "a:blipFill/a:blip")
                if picture is not None:
                    source = self.source(part, picture.get(qn("r:embed")), background)
                    if source is not None:
                        return f' style="background:url(&quot;{html.escape(source)}&quot;) center/cover"'
                color = fill_color(self.theme, properties, None)
            elif reference is not None:
                color = self.theme.color(reference)
            else:
                continue
            if color is None or color.css() == "#FFFFFF":
                return ""
            return f' style="background:{color.css()}"'
        return ""


_FORMATS = {"PNG": ".png", "JPEG": ".jpg", "GIF": ".gif", "WEBP": ".webp"}
"""The extension each picture format browsers show is given, by Pillow's name for it."""

_CONVERTED = {"TIFF", "BMP"}
"""Picture formats most browsers do not show, which are converted to PNG."""


class _Media:
    """The folder the pictures and movies on a page are copied into."""

    def __init__(
        self,
        folder: pathlib.Path,
        substitutes: "dict[str, str]",
        reencode: "Reencode | None",
        relink: "dict[str, str]",
    ):
        self.folder = folder
        self.substitutes = substitutes
        self.reencode = reencode
        self.relink = [(_slashes(old).lower(), _slashes(new)) for old, new in relink.items()]
        self.written: dict[str, str] = {}
        """The address each file was given, by the part or path it came from."""
        self.names: set[str] = set()
        """The names already taken in the folder."""

    def add(self, part: Part, movie: bool = False) -> str:
        """The address of a picture or movie in the deck, copying it the first time."""
        name = posixpath.basename(str(part.partname))
        if name in self.substitutes:
            return self.substitutes[name]
        return self._store(str(part.partname), name, part.blob, movie)

    def add_file(self, path: pathlib.Path, movie: bool = False) -> str:
        """The address of a picture or movie the deck links to, copying it the first time."""
        return self._store(str(path.resolve()), path.name, path.read_bytes(), movie)

    def locate(self, target: str) -> "pathlib.Path | None":
        """Where a file a deck links to is on this computer, if anywhere."""
        path = _slashes(target)
        for prefix in ("file:///", "file://"):
            path = path.removeprefix(prefix)
        for old, new in self.relink:
            if path.lower().startswith(old):
                path = new + path[len(old):]
                break
        candidate = pathlib.Path(path)
        return candidate if candidate.is_file() else None

    def _store(self, key: str, name: str, data: bytes, movie: bool) -> str:
        if key in self.written:
            return self.written[key]
        self.folder.mkdir(parents=True, exist_ok=True)
        name, data = self._movie(name, data) if movie else self._picture(name, data)
        name = self._unique(name)
        (self.folder / name).write_bytes(data)
        address = f"media/{name}"
        self.written[key] = address
        return address

    def _unique(self, name: str) -> str:
        """A name not yet taken in the folder: two linked files can share one."""
        stem, suffix = posixpath.splitext(name)
        candidate = name
        number = 1
        while candidate in self.names:
            number += 1
            candidate = f"{stem}-{number}{suffix}"
        self.names.add(candidate)
        return candidate

    def _movie(self, name: str, data: bytes) -> "tuple[str, bytes]":
        """A movie, re-encoded if asked and if that makes it smaller."""
        if self.reencode is None:
            return name, data
        encoded = self._encode(name, data)
        if len(encoded) < len(data) or not name.lower().endswith(".mp4"):
            return posixpath.splitext(name)[0] + ".mp4", encoded
        return name, data

    def _picture(self, name: str, data: bytes) -> "tuple[str, bytes]":
        """
        A picture in a form browsers show: given the extension its contents
        call for, converted to PNG if it is a TIFF or BMP, and turned into a
        movie if it is an animated GIF and movies are being re-encoded, when
        the movie is the smaller.
        """
        stem, suffix = posixpath.splitext(name)
        if suffix.lower() == ".svg" or data.lstrip()[:1] == b"<":
            return name, data
        try:
            image = PIL.Image.open(io.BytesIO(data))
        except PIL.UnidentifiedImageError:
            return name, data
        kind = image.format or ""
        if kind == "GIF" and getattr(image, "n_frames", 1) > 1 and self.reencode is not None:
            encoded = self._encode(name, data)
            if len(encoded) < len(data):
                return f"{stem}.mp4", encoded
        if kind in _FORMATS:
            extensions = {_FORMATS[kind]} | ({".jpeg"} if kind == "JPEG" else set())
            if suffix.lower() not in extensions:
                return stem + _FORMATS[kind], data
            return name, data
        if kind in _CONVERTED:
            buffer = io.BytesIO()
            if image.mode not in ("1", "L", "LA", "P", "RGB", "RGBA"):
                image = image.convert("RGBA")
            image.save(buffer, "PNG")
            return f"{stem}.png", buffer.getvalue()
        return name, data

    def _encode(self, name: str, data: bytes) -> bytes:
        """A movie or animated GIF re-encoded as H.264, for the web."""
        assert self.reencode is not None
        with tempfile.TemporaryDirectory() as scratch:
            source = pathlib.Path(scratch) / name
            source.write_bytes(data)
            encoded = pathlib.Path(scratch) / "encoded.mp4"
            width = self.reencode.max_width
            filters = [f"scale='min({width},trunc(iw/2)*2)':-2"]
            timing = _gif_timing(data)
            if timing is not None:
                # A movie at a steady frame rate holds each frame until the
                # next one, and the last has no next one, so a GIF that ends
                # by holding a frame, like one that blinks between two
                # pictures, would lose that hold. So the last frame is copied
                # on past its end, the frames are laid out at a steady rate,
                # and the movie is cut to the length the GIF plays for.
                total, hold, rate = timing
                filters[:0] = [
                    f"tpad=stop_mode=clone:stop_duration={hold}",
                    f"fps={rate}",
                    f"trim=duration={total:.3f}",
                ]
            command = [
                self.reencode.ffmpeg, "-y", "-loglevel", "error",
                "-i", str(source),
                "-an",
                "-c:v", "libx264", "-preset", "slow", "-crf", str(self.reencode.crf),
                "-pix_fmt", "yuv420p",
                "-vf", ",".join(filters),
                "-movflags", "+faststart",
                str(encoded),
            ]
            subprocess.run(command, check=True)
            return encoded.read_bytes()


def _gif_timing(data: bytes) -> "tuple[float, float, int] | None":
    """
    How long an animated GIF plays and how long it holds its last frame, in
    seconds, and the frame rate that keeps its quickest frame, or `None` if
    the data is not an animated GIF. Delays of a hundredth of a second or
    less are shown for a tenth, as browsers and ffmpeg both do.
    """
    if data[:6] not in (b"GIF87a", b"GIF89a"):
        return None
    delays = []
    with PIL.Image.open(io.BytesIO(data)) as image:
        for frame in range(getattr(image, "n_frames", 1)):
            image.seek(frame)
            delay = image.info.get("duration", 0)
            delays.append((delay if delay > 10 else 100) / 1000)
    if len(delays) < 2:
        return None
    rate = min(60, max(1, round(1 / min(delays))))
    return sum(delays), delays[-1], rate


def _slashes(path: str) -> str:
    return path.replace("\\", "/")


def _walk(
    tree: "etree._Element | None",
    transform: "_Transform",
) -> "Iterator[tuple[etree._Element, _Transform]]":
    """Every drawable shape in a shape tree, with how to place it on the slide."""
    for child in tree if tree is not None else []:
        kind = local(child)
        if kind in ("sp", "pic", "cxnSp", "graphicFrame"):
            yield child, transform
        elif kind == "grpSp":
            frame = find(child, "p:grpSpPr/a:xfrm")
            offset, extent = find(frame, "a:off"), find(frame, "a:ext")
            child_offset, child_extent = find(frame, "a:chOff"), find(frame, "a:chExt")
            gx, gy = integer(offset, "x"), integer(offset, "y")
            cx, cy = integer(child_offset, "x"), integer(child_offset, "y")
            sx = integer(extent, "cx", 1) / (integer(child_extent, "cx", 1) or 1)
            sy = integer(extent, "cy", 1) / (integer(child_extent, "cy", 1) or 1)

            def inner(box: _Box, gx=gx, gy=gy, cx=cx, cy=cy, sx=sx, sy=sy, outer=transform) -> _Box:
                return outer(
                    dataclasses.replace(
                        box,
                        x=gx + (box.x - cx) * sx,
                        y=gy + (box.y - cy) * sy,
                        width=box.width * sx,
                        height=box.height * sy,
                    )
                )

            yield from _walk(child, inner)
        elif kind == "AlternateContent":
            yield from _walk(_alternative(child), transform)


_COMPATIBILITY = "http://schemas.openxmlformats.org/markup-compatibility/2006"

_UNDERSTOOD = {
    # Office 2010 drawing, which is what equations in text need.
    "http://schemas.microsoft.com/office/drawing/2010/main",
}
"""The newer namespaces whose content this can draw itself."""


def _alternative(content: etree._Element) -> "etree._Element | None":
    """
    The version of some newer content to draw: the newer one, when everything
    it needs is understood, and otherwise the older stand-in that comes with
    it. A text box with an equation in it is the usual case: its stand-in is
    a picture of the text, laid over an empty box.
    """
    for choice in content.findall(f"{{{_COMPATIBILITY}}}Choice"):
        required = choice.get("Requires", "").split()
        if required and all(choice.nsmap.get(prefix) in _UNDERSTOOD for prefix in required):
            return choice
    return content.find(f"{{{_COMPATIBILITY}}}Fallback")


def _box(element: etree._Element, chain: "list[etree._Element]") -> "_Box | None":
    """Where a shape is, taking a placeholder's position from its layout if it has none."""
    for candidate in [element] + chain:
        frame = find(candidate, "p:spPr/a:xfrm")
        if frame is None:
            frame = find(candidate, "p:xfrm")
        offset, extent = find(frame, "a:off"), find(frame, "a:ext")
        if offset is None or extent is None:
            continue
        own = find(element, "p:spPr/a:xfrm")
        return _Box(
            x=integer(offset, "x"),
            y=integer(offset, "y"),
            width=integer(extent, "cx"),
            height=integer(extent, "cy"),
            rotation=integer(own, "rot") / 60000,
            flip_h=(own is not None and own.get("flipH") in ("1", "true")),
            flip_v=(own is not None and own.get("flipV") in ("1", "true")),
        )
    return None


_MOVIES = (".mp4", ".m4v", ".mov", ".webm")
"""The file types a picture can be swapped for and still play, as a movie."""


def _attribute(element: "etree._Element | None", attribute: str) -> "str | None":
    """A namespaced attribute like ``r:embed`` of an element that may be missing."""
    return None if element is None else element.get(qn(attribute))


def _placeholder(element: etree._Element) -> "etree._Element | None":
    for path in ("p:nvSpPr/p:nvPr/p:ph", "p:nvPicPr/p:nvPr/p:ph", "p:nvGraphicFramePr/p:nvPr/p:ph"):
        found = find(element, path)
        if found is not None:
            return found
    return None


def _match(owner: etree._Element, kind: str, index: "str | None") -> "etree._Element | None":
    """The placeholder on a layout or master that a slide's placeholder is based on."""
    tree = find(owner, "p:cSld/p:spTree")
    candidates = []
    for shape in tree if tree is not None else []:
        placeholder = _placeholder(shape)
        if placeholder is not None:
            candidates.append((shape, placeholder))
    if index is not None and index != "0":
        for shape, placeholder in candidates:
            if placeholder.get("idx", "0") == index:
                return shape
    for shape, placeholder in candidates:
        if placeholder.get("type", "obj") == kind:
            return shape
    if kind in ("ctrTitle", "title"):
        for shape, placeholder in candidates:
            if placeholder.get("type") in ("ctrTitle", "title"):
                return shape
    return None


def _notes(slide: Slide) -> str:
    """A slide's speaker notes as paragraphs of HTML, or nothing if it has none."""
    if not slide.has_notes_slide:
        return ""
    notes_slide = slide.notes_slide
    tree = find(notes_slide.element, "p:cSld/p:spTree")
    paragraphs = []
    for shape in tree if tree is not None else []:
        placeholder = _placeholder(shape)
        # The notes are the body placeholder; the others on a notes page are
        # the picture of the slide, its number, and the like.
        if placeholder is None or placeholder.get("type") != "body":
            continue
        for paragraph in findall(shape, "p:txBody/a:p"):
            if not "".join(t.text or "" for t in findall(paragraph, ".//a:t")).strip():
                continue
            level = integer(find(paragraph, "a:pPr"), "lvl")
            indent = f' style="margin-left:{1.5 * level:g}em"' if level else ""
            paragraphs.append(f"<p{indent}>{_note(paragraph, notes_slide.part)}</p>")
    if not paragraphs:
        return ""
    return f'<div class="notes">{"".join(paragraphs)}</div>'


def _note(paragraph: etree._Element, part: Part) -> str:
    """One paragraph of speaker notes, keeping its emphasis and links."""
    pieces = []
    for element in paragraph:
        kind = local(element)
        if kind == "br":
            pieces.append("<br>")
            continue
        if kind == "m":
            pieces.append(mathml(element, lambda message: None))
            continue
        if kind not in ("r", "fld"):
            continue
        text = html.escape(element.findtext(qn("a:t")) or "", quote=False)
        properties = find(element, "a:rPr")
        if properties is not None and text:
            if properties.get("b") in ("1", "true"):
                text = f"<strong>{text}</strong>"
            if properties.get("i") in ("1", "true"):
                text = f"<em>{text}</em>"
            if (properties.get("u") or "none") != "none":
                text = f"<u>{text}</u>"
            link = find(properties, "a:hlinkClick")
            address = _hyperlink(part, link.get(qn("r:id"), "")) if link is not None else None
            if address:
                text = f'<a href="{html.escape(address)}">{text}</a>'
        pieces.append(text)
    return "".join(pieces)


def _first_title(slides: "list[Slide]") -> "str | None":
    """The text of the title on the first slide, if it has one."""
    for slide in slides[:1]:
        tree = find(slide.element, "p:cSld/p:spTree")
        for shape in tree if tree is not None else []:
            placeholder = _placeholder(shape)
            if placeholder is not None and placeholder.get("type") in ("title", "ctrTitle"):
                text = " ".join(" ".join(t.text or "" for t in findall(shape, ".//a:t")).split())
                if text:
                    return text
    return None


def _name(element: etree._Element) -> str:
    for properties in element.iter(qn("p:cNvPr")):
        return properties.get("name", "a shape")
    return "a shape"


def _description(element: etree._Element) -> str:
    """A picture's alt text, without the caption Office writes for it itself."""
    for properties in element.iter(qn("p:cNvPr")):
        text = properties.get("descr", "")
        lines = [
            line for line in text.splitlines()
            if line.strip() and line.strip() != "Description automatically generated"
        ]
        return " ".join(lines)
    return ""


def _hyperlink(part: Part, rid: str) -> "str | None":
    if rid in part.rels and part.rels[rid].is_external:
        return part.rels[rid].target_ref
    return None


def _scheme(name: str) -> etree._Element:
    fill = etree.Element(qn("a:solidFill"), nsmap={"a": NS["a"]})
    etree.SubElement(fill, qn("a:schemeClr"), val=name)
    return fill


def _fixed(value: float) -> str:
    """A number to four decimal places, never as negative zero."""
    text = f"{value:.4f}"
    return "0.0000" if text == "-0.0000" else text


def _number(value: float) -> str:
    return f"{value:.3f}".rstrip("0").rstrip(".") if value % 1 else f"{value:.0f}"


def _integer_or(value: "str | None", default: int) -> int:
    return int(value) if value is not None else default

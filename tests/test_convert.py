import io
import pathlib
import shutil
import subprocess
import zipfile
from typing import TypeVar, cast

import pytest
from lxml import etree, html
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.presentation import Presentation as Deck
from pptx.shapes.placeholder import SlidePlaceholder
from pptx.util import Emu, Pt

import pptx_to_html

A = "http://schemas.openxmlformats.org/drawingml/2006/main"

WIDTH = 12192000
HEIGHT = 6858000

LAYOUT_TITLE = 0
LAYOUT_CONTENT = 1
LAYOUT_BLANK = 6

ORIGIN = Emu(0)

T = TypeVar("T")


def one(found: "T | None") -> T:
    """Something a test expects to find, which fails the test if it is not there."""
    assert found is not None
    return found


def png(width: int = 40, height: int = 20) -> io.BytesIO:
    buffer = io.BytesIO()
    Image.new("RGB", (width, height), (200, 30, 30)).save(buffer, "PNG")
    buffer.seek(0)
    return buffer


def deck() -> Deck:
    presentation = Presentation()
    presentation.slide_width = Emu(WIDTH)
    presentation.slide_height = Emu(HEIGHT)
    return presentation


def convert(presentation: Deck, folder: pathlib.Path, **kwargs) -> html.HtmlElement:
    path = folder / "deck.pptx"
    presentation.save(str(path))
    index = pptx_to_html.convert(path, folder / "out", **kwargs)
    return html.parse(str(index)).getroot()


def style(element: html.HtmlElement) -> "dict[str, str]":
    pairs = [p.split(":", 1) for p in element.get("style", "").split(";") if ":" in p]
    return {k.strip(): v.strip() for k, v in pairs}


def slide(page: html.HtmlElement, number: int) -> html.HtmlElement:
    return page.get_element_by_id(f"slide-{number}")


def test_title_slide(tmp_path):
    presentation = deck()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_TITLE])
    one(title_slide.shapes.title).text_frame.text = "A Talk About the Sun"
    subtitle = title_slide.placeholders[1]
    assert isinstance(subtitle, SlidePlaceholder)
    subtitle.text_frame.text = "Someone"

    page = convert(presentation, tmp_path)

    assert page.findtext(".//title") == "A Talk About the Sun"
    assert page.findtext(".//h1") == "A Talk About the Sun"
    title = slide(page, 1).xpath(".//p[normalize-space()='A Talk About the Sun']")[0]
    # Nothing on the default template's title slide sets a size, so the title
    # takes the master's 44 points, on a slide 540 points high.
    assert style(title)["font-size"] == f"{100 * 44 / 540:.4f}cqh"
    assert style(title)["text-align"] == "center"
    # The default theme sets headings in the same typeface as everything else.
    assert "font-family" not in style(title)


def test_subtitle(tmp_path):
    presentation = deck()
    presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_BLANK])
    page = convert(presentation, tmp_path, title="Talk", subtitle="A meeting · 2026")
    assert page.findtext(".//header/p") == "A meeting · 2026"
    assert page.findtext(".//div[@class='number']") == "1 / 1"


def test_text_box(tmp_path):
    presentation = deck()
    blank = presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_BLANK])
    box = blank.shapes.add_textbox(Emu(WIDTH // 4), Emu(HEIGHT // 2), Emu(WIDTH // 2), Emu(HEIGHT // 4))
    box.text_frame.word_wrap = True
    paragraph = box.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = "Hello"
    run.font.size = Pt(32)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x12, 0x34, 0x56)

    page = convert(presentation, tmp_path)

    piece = slide(page, 1).find_class("text")[0]
    assert style(piece)["left"] == "25.0000%"
    assert style(piece)["top"] == "50.0000%"
    assert style(piece)["width"] == "50.0000%"
    assert style(piece)["height"] == "25.0000%"
    assert style(piece)["padding"] == "0.667cqh 1.333cqh 0.667cqh 1.333cqh"
    p = one(piece.find("p"))
    assert style(p)["font-size"] == f"{100 * 32 / 540:.4f}cqh"
    assert style(p)["color"] == "#123456"
    assert style(p)["font-weight"] == "bold"
    assert p.text_content() == "Hello"


def test_text_box_without_wrapping(tmp_path):
    presentation = deck()
    blank = presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_BLANK])
    box = blank.shapes.add_textbox(ORIGIN, ORIGIN, Emu(WIDTH // 8), Emu(HEIGHT // 8))
    box.text_frame.word_wrap = False
    box.text_frame.text = "A label longer than its box"

    page = convert(presentation, tmp_path, title="Talk")

    piece = slide(page, 1).find_class("text")[0]
    assert style(piece)["white-space"] == "nowrap"
    assert style(piece)["width"] == "max-content"


def test_bullets(tmp_path):
    presentation = deck()
    content = presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_CONTENT])
    one(content.shapes.title).text_frame.text = "Points"
    placeholder = content.placeholders[1]
    assert isinstance(placeholder, SlidePlaceholder)
    body = placeholder.text_frame
    body.text = "First"
    second = body.add_paragraph()
    second.text = "Second, indented"
    second.level = 1

    page = convert(presentation, tmp_path)

    bullets = slide(page, 1).find_class("bullet")
    assert [b.text_content() for b in bullets] == ["First", "Second, indented"]
    first, indented = (style(b) for b in bullets)
    assert first["--bullet"] == "'•'"
    assert float(indented["padding-left"].removesuffix("cqh")) > float(
        first.get("padding-left", "0cqh").removesuffix("cqh")
    )
    assert float(indented["font-size"].removesuffix("cqh")) < float(first["font-size"].removesuffix("cqh"))


def test_cropped_picture(tmp_path):
    presentation = deck()
    blank = presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_BLANK])
    picture = blank.shapes.add_picture(png(), ORIGIN, ORIGIN, Emu(WIDTH // 2), Emu(HEIGHT // 2))
    picture.crop_left = 0.2

    page = convert(presentation, tmp_path)

    wrapper = slide(page, 1).find_class("crop")[0]
    image = one(wrapper.find("img"))
    assert style(wrapper)["width"] == "50.0000%"
    assert style(image)["width"] == "125.0000%"
    assert style(image)["left"] == "-25.0000%"
    assert style(image)["top"] == "0.0000%"
    assert (tmp_path / "out" / image.get("src", "")).is_file()


def test_octagon(tmp_path):
    presentation = deck()
    blank = presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_BLANK])
    octagon = blank.shapes.add_shape(MSO_SHAPE.OCTAGON, ORIGIN, ORIGIN, Emu(WIDTH), Emu(HEIGHT // 2))
    octagon.fill.solid()
    octagon.fill.fore_color.rgb = RGBColor(0, 0, 0)
    octagon.line.fill.background()

    page = convert(presentation, tmp_path)

    piece = slide(page, 1).find_class("piece")[0]
    assert style(piece)["background"] == "#000000"
    assert style(piece)["clip-path"].startswith("polygon(")
    assert "border" not in style(piece)


def test_movie(tmp_path):
    presentation = deck()
    blank = presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_BLANK])
    movie = tmp_path / "movie.mp4"
    movie.write_bytes(b"not really a movie")
    blank.shapes.add_movie(str(movie), ORIGIN, ORIGIN, Emu(WIDTH), Emu(HEIGHT), poster_frame_image=png(), mime_type="video/mp4")

    page = convert(presentation, tmp_path)

    video = one(slide(page, 1).find(".//video"))
    assert video.get("src", "").startswith("media/")
    assert video.get("poster", "").startswith("media/")
    for name in ("loop", "muted", "playsinline", "controls"):
        assert name in video.attrib
    assert (tmp_path / "out" / video.get("src", "")).read_bytes() == b"not really a movie"


def test_substitutes(tmp_path):
    presentation = deck()
    blank = presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_BLANK])
    blank.shapes.add_picture(png(), ORIGIN, ORIGIN, Emu(WIDTH // 2), Emu(HEIGHT // 2))
    blank.shapes.add_picture(png(41, 20), Emu(WIDTH // 2), ORIGIN, Emu(WIDTH // 2), Emu(HEIGHT // 2))

    path = tmp_path / "deck.pptx"
    presentation.save(str(path))
    names = sorted(n for n in zipfile.ZipFile(path).namelist() if n.startswith("ppt/media/"))
    first, second = (n.rsplit("/", 1)[-1] for n in names)
    index = pptx_to_html.convert(
        path,
        tmp_path / "out",
        media={first: "https://example.com/sharper.png", second: "../figures/animation.mp4"},
    )
    page = html.parse(str(index)).getroot()

    assert one(slide(page, 1).find(".//img")).get("src") == "https://example.com/sharper.png"
    # A picture swapped for a movie plays as one.
    assert one(slide(page, 1).find(".//video")).get("src") == "../figures/animation.mp4"


def test_hidden_slides(tmp_path):
    presentation = deck()
    for _ in range(3):
        presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_BLANK])
    presentation.slides[1]._element.set("show", "0")

    page = convert(presentation, tmp_path, title="Talk")
    assert len(page.find_class("slide")) == 2

    page = convert(presentation, tmp_path, title="Talk", include_hidden=True)
    assert len(page.find_class("slide")) == 3


def test_cli(tmp_path, capsys):
    from pptx_to_html.__main__ import main

    presentation = deck()
    presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_BLANK])
    path = tmp_path / "deck.pptx"
    presentation.save(str(path))

    assert main([str(path), str(tmp_path / "out"), "--title", "Talk", "--notes"]) == 0
    assert (tmp_path / "out" / "index.html").is_file()
    assert "index.html" in capsys.readouterr().out


@pytest.mark.skipif(shutil.which("ffmpeg") is None, reason="needs ffmpeg")
def test_reencode(tmp_path):
    movie = tmp_path / "movie.mp4"
    subprocess.run(
        [
            "ffmpeg", "-y", "-loglevel", "error",
            "-f", "lavfi", "-i", "testsrc=size=320x240:rate=10:duration=1",
            "-c:v", "mpeg4", "-q:v", "1", str(movie),
        ],
        check=True,
    )
    presentation = deck()
    blank = presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_BLANK])
    blank.shapes.add_movie(str(movie), ORIGIN, ORIGIN, Emu(WIDTH), Emu(HEIGHT), poster_frame_image=png(), mime_type="video/mp4")

    page = convert(presentation, tmp_path, title="Talk", reencode=pptx_to_html.Reencode())

    video = one(slide(page, 1).find(".//video"))
    encoded = tmp_path / "out" / video.get("src", "")
    assert encoded.suffix == ".mp4"
    assert encoded.stat().st_size <= movie.stat().st_size


def test_group(tmp_path):
    presentation = deck()
    blank = presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_BLANK])
    group = blank.shapes.add_group_shape()
    square = group.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(WIDTH // 4), Emu(HEIGHT // 4), Emu(WIDTH // 2), Emu(HEIGHT // 2)
    )
    square.fill.solid()
    square.fill.fore_color.rgb = RGBColor(0, 0x80, 0)

    page = convert(presentation, tmp_path, title="Talk")

    piece = slide(page, 1).find_class("piece")[0]
    assert style(piece)["left"] == "25.0000%"
    assert style(piece)["width"] == "50.0000%"
    assert style(piece)["background"] == "#008000"


def test_line(tmp_path):
    presentation = deck()
    blank = presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_BLANK])
    line = blank.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, ORIGIN, Emu(HEIGHT), Emu(WIDTH), ORIGIN)
    line.line.color.rgb = RGBColor(0xFF, 0, 0)
    line.line.width = Pt(2)

    page = convert(presentation, tmp_path, title="Talk")

    drawn = one(slide(page, 1).find(".//path"))
    assert drawn.get("stroke") == "#FF0000"
    assert drawn.get("stroke-width") == str(Pt(2))
    # Drawn from the bottom left corner to the top right one.
    assert drawn.get("d") == f"M0,{HEIGHT} L{WIDTH},0"


def test_elbow_connector(tmp_path):
    # The elbow on a slide of the 2022 LMSAL colloquium that comes out of the
    # left of one box and goes down into another: turned half a circle and
    # flipped, its corner is at the top left.
    presentation = deck()
    blank = presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_BLANK])
    elbow = blank.shapes.add_connector(MSO_CONNECTOR.ELBOW, ORIGIN, ORIGIN, Emu(100), Emu(100))
    elbow.line.color.rgb = RGBColor(0, 0, 0)
    element = cast(etree._Element, elbow._element)
    # python-pptx's elbow bends twice; the one on the slide bends once.
    geometry = element.find(f".//{{{A}}}prstGeom")
    assert geometry is not None
    geometry.set("prst", "bentConnector2")
    frame = element.find(f".//{{{A}}}xfrm")
    assert frame is not None
    frame.attrib.clear()
    frame.set("rot", "10800000")
    frame.set("flipV", "1")
    offset, extent = frame.find(f"{{{A}}}off"), frame.find(f"{{{A}}}ext")
    assert offset is not None and extent is not None
    offset.set("x", "8097918")
    offset.set("y", "3579853")
    extent.set("cx", "422932")
    extent.set("cy", "1213337")

    page = convert(presentation, tmp_path, title="Talk")

    drawn = one(slide(page, 1).find(".//path"))
    assert drawn.get("d") == "M8520850,3579853 L8097918,3579853 L8097918,4793190"


def test_background(tmp_path):
    presentation = deck()
    blank = presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_BLANK])
    blank.background.fill.solid()
    blank.background.fill.fore_color.rgb = RGBColor(0x10, 0x20, 0x30)

    page = convert(presentation, tmp_path, title="Talk")

    stage = slide(page, 1).find_class("stage")[0]
    assert style(stage)["background"] == "#102030"


def test_notes(tmp_path):
    presentation = deck()
    first = presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_BLANK])
    frame = one(first.notes_slide.notes_text_frame)
    frame.text = "Hello everyone."
    said = frame.add_paragraph()
    run = said.add_run()
    run.text = "This part matters."
    run.font.bold = True
    presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_BLANK])

    # Notes are left out unless asked for.
    page = convert(presentation, tmp_path, title="Talk")
    assert not page.find_class("notes")
    assert page.get_element_by_id("notes", None) is None

    page = convert(presentation, tmp_path, title="Talk", notes=True)
    remarks = slide(page, 1).find_class("notes")
    assert len(remarks) == 1
    paragraphs = remarks[0].findall("p")
    assert [p.text_content() for p in paragraphs] == ["Hello everyone.", "This part matters."]
    assert one(paragraphs[1].find("strong")).text == "This part matters."
    # A slide without notes gets nothing under it.
    assert not slide(page, 2).find_class("notes")
    assert one(page.get_element_by_id("notes")).text == "Hide notes"


def test_autofit(tmp_path):
    presentation = deck()
    blank = presentation.slides.add_slide(presentation.slide_layouts[LAYOUT_BLANK])
    shrinking = blank.shapes.add_textbox(ORIGIN, ORIGIN, Emu(WIDTH // 2), Emu(HEIGHT // 8))
    shrinking.text_frame.word_wrap = True
    run = shrinking.text_frame.paragraphs[0].add_run()
    run.text = "Shrunk when it overflows"
    run.font.size = Pt(27)
    body = cast(etree._Element, shrinking.text_frame._txBody).find(f"{{{A}}}bodyPr")
    assert body is not None
    for child in list(body):
        body.remove(child)
    # Saved with the shrink PowerPoint worked out, which it draws the box with.
    etree.SubElement(body, f"{{{A}}}normAutofit", fontScale="50000", lnSpcReduction="20000")
    plain = blank.shapes.add_textbox(ORIGIN, Emu(HEIGHT // 2), Emu(WIDTH // 2), Emu(HEIGHT // 8))
    plain.text_frame.word_wrap = True
    plain.text_frame.text = "Left to overflow"

    page = convert(presentation, tmp_path, title="Talk")

    boxes = slide(page, 1).find_class("text")
    fitted = [b for b in boxes if "fit" in b.get("class", "").split()]
    assert [b.text_content() for b in fitted] == ["Shrunk when it overflows"]
    p = one(fitted[0].find("p"))
    assert style(p)["font-size"] == f"{100 * 27 * 0.5 / 540:.4f}cqh"
    assert style(p)["line-height"] == f"{1.2 * (1 - 0.2):.3f}"

import io
import pathlib
import shutil
import subprocess
import zipfile
from typing import cast

import pytest
from lxml import etree, html
from PIL import Image
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.presentation import Presentation as Deck
from pptx.slide import Slide
from pptx.util import Emu

import pptx_to_html

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

WIDTH = 12192000
HEIGHT = 6858000


def picture(kind: str = "PNG", **options) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (40, 20), (200, 30, 30)).save(buffer, kind, **options)
    return buffer.getvalue()


def deck() -> tuple[Deck, Slide]:
    presentation = Presentation()
    presentation.slide_width = Emu(WIDTH)
    presentation.slide_height = Emu(HEIGHT)
    return presentation, presentation.slides.add_slide(presentation.slide_layouts[6])


def link(slide: Slide, target: str) -> None:
    """A picture on the slide that links to ``target`` instead of embedding it."""
    shape = slide.shapes.add_picture(io.BytesIO(picture()), Emu(0), Emu(0), Emu(WIDTH // 4), Emu(HEIGHT // 4))
    rid = slide.part.relate_to(target, RT.IMAGE, is_external=True)
    blip = cast(etree._Element, shape._element).find(f".//{{{A}}}blip")
    assert blip is not None
    del blip.attrib[f"{{{R}}}embed"]
    blip.set(f"{{{R}}}link", rid)


def convert(presentation: Deck, folder: pathlib.Path, **kwargs) -> html.HtmlElement:
    path = folder / "deck.pptx"
    presentation.save(str(path))
    return html.parse(str(pptx_to_html.convert(path, folder / "out", title="Talk", **kwargs))).getroot()


def sources(page: html.HtmlElement) -> list[str]:
    return [element.get("src", "") for element in page.iter("img", "video")]


def test_linked_file_where_the_deck_says(tmp_path):
    there = tmp_path / "figures" / "plot.png"
    there.parent.mkdir()
    there.write_bytes(picture())
    presentation, slide = deck()
    link(slide, there.as_uri())

    page = convert(presentation, tmp_path)

    assert sources(page) == ["media/plot.png"]
    assert (tmp_path / "out" / "media" / "plot.png").read_bytes() == there.read_bytes()


def test_relinked_files(tmp_path):
    moved = tmp_path / "moved"
    for folder in ("first", "second"):
        (moved / folder).mkdir(parents=True)
        (moved / folder / "zoom.png").write_bytes(picture())
    presentation, slide = deck()
    link(slide, "file:///C:/Users/someone/talk/first/zoom.png")
    link(slide, "file:///C:/Users/someone/talk/second/zoom.png")

    page = convert(presentation, tmp_path, relink={"C:\\Users\\someone\\talk": str(moved)})

    # Two files that share a name both make it onto the page.
    assert sources(page) == ["media/zoom.png", "media/zoom-2.png"]


def test_lost_link_warns(tmp_path):
    presentation, slide = deck()
    link(slide, "file:///C:/nowhere/plot.png")

    with pytest.warns(pptx_to_html.ConversionWarning, match="--relink"):
        page = convert(presentation, tmp_path)

    assert sources(page) == []


def test_tiff_becomes_png(tmp_path):
    presentation, slide = deck()
    slide.shapes.add_picture(io.BytesIO(picture("TIFF")), Emu(0), Emu(0), Emu(WIDTH // 4), Emu(HEIGHT // 4))

    page = convert(presentation, tmp_path)

    (source,) = sources(page)
    assert source.endswith(".png")
    with Image.open(tmp_path / "out" / source) as image:
        assert image.format == "PNG"


def test_picture_gets_its_extension(tmp_path):
    presentation, slide = deck()
    slide.shapes.add_picture(io.BytesIO(picture()), Emu(0), Emu(0), Emu(WIDTH // 4), Emu(HEIGHT // 4))
    path = tmp_path / "deck.pptx"
    presentation.save(str(path))
    # Save the picture inside the deck under a name that hides what it is,
    # as PowerPoint sometimes does.
    renamed = tmp_path / "renamed.pptx"
    with zipfile.ZipFile(path) as original, zipfile.ZipFile(renamed, "w") as copy:
        for info in original.infolist():
            data = original.read(info.filename)
            if info.filename.startswith("ppt/media/"):
                info.filename = info.filename.rsplit(".", 1)[0] + ".bin"
            elif info.filename.endswith(".rels") or info.filename == "[Content_Types].xml":
                data = data.replace(b".png", b".bin").replace(b'Extension="png"', b'Extension="bin"')
            copy.writestr(info, data)

    index = pptx_to_html.convert(renamed, tmp_path / "out", title="Talk")

    (source,) = sources(html.parse(str(index)).getroot())
    assert source.endswith(".png")


@pytest.mark.skipif(shutil.which("ffmpeg") is None, reason="needs ffmpeg")
def test_animated_gif_becomes_movie(tmp_path):
    # A gradient sliding across the frame: easy for a movie, hard for a GIF.
    frames = []
    for shift in range(40):
        frame = Image.new("RGB", (200, 200))
        frame.putdata([((x + 5 * shift) % 256, y, (x + y) % 256) for y in range(200) for x in range(200)])
        frames.append(frame)
    buffer = io.BytesIO()
    frames[0].save(buffer, "GIF", save_all=True, append_images=frames[1:], duration=50, loop=0)
    presentation, slide = deck()
    buffer.seek(0)
    slide.shapes.add_picture(buffer, Emu(0), Emu(0), Emu(WIDTH // 2), Emu(HEIGHT // 2))

    page = convert(presentation, tmp_path, reencode=pptx_to_html.Reencode())

    video = page.find(".//video")
    assert video is not None
    assert video.get("src", "").endswith(".mp4")
    assert page.find(".//img") is None


@pytest.mark.skipif(
    shutil.which("ffmpeg") is None or shutil.which("ffprobe") is None,
    reason="needs ffmpeg and ffprobe",
)
def test_gif_keeps_its_last_frame(tmp_path):
    # An animation that ends by holding its last frame for two seconds.
    frames = []
    for shift in range(40):
        frame = Image.new("RGB", (200, 200))
        frame.putdata([((x + 5 * shift) % 256, y, (x + y) % 256) for y in range(200) for x in range(200)])
        frames.append(frame)
    buffer = io.BytesIO()
    durations = [50] * 39 + [2000]
    frames[0].save(buffer, "GIF", save_all=True, append_images=frames[1:], duration=durations, loop=0)
    presentation, slide = deck()
    buffer.seek(0)
    slide.shapes.add_picture(buffer, Emu(0), Emu(0), Emu(WIDTH // 2), Emu(HEIGHT // 2))

    page = convert(presentation, tmp_path, reencode=pptx_to_html.Reencode())

    video = page.find(".//video")
    assert video is not None
    probe = subprocess.run(
        ["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0",
         str(tmp_path / "out" / video.get("src", ""))],
        capture_output=True, text=True, check=True,
    )
    assert float(probe.stdout) == pytest.approx(39 * 0.05 + 2, abs=0.15)

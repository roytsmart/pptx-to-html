import io
import pathlib
from typing import cast

import pytest
from lxml import etree, html
from PIL import Image
from pptx import Presentation
from pptx.presentation import Presentation as Deck
from pptx.shapes.autoshape import Shape
from pptx.slide import Slide
from pptx.util import Emu

import pptx_to_html
from pptx_to_html._fonts import font_faces, font_stack
from pptx_to_html._math import mathml

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
A14 = "http://schemas.microsoft.com/office/drawing/2010/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"

WIDTH = 12192000
HEIGHT = 6858000


def equation(body: str) -> etree._Element:
    """An ``a14:m`` element around some Office Math."""
    return etree.fromstring(f'<a14:m xmlns:a14="{A14}" xmlns:m="{M}"><m:oMath>{body}</m:oMath></a14:m>')


def convert(element: etree._Element) -> "tuple[str, list[str]]":
    warnings: list[str] = []
    return mathml(element, warnings.append), warnings


def test_runs():
    result, warnings = convert(equation("<m:r><m:t>𝑥+12.5</m:t></m:r>"))
    assert result == "<math><mrow><mi>𝑥</mi><mo>+</mo><mn>12.5</mn></mrow></math>"
    assert not warnings


def test_upright_text():
    result, _ = convert(equation("<m:r><m:rPr><m:nor/></m:rPr><m:t>VMR</m:t></m:r>"))
    assert "<mi>VMR</mi>" in result


def test_fraction():
    result, _ = convert(equation("<m:f><m:num><m:r><m:t>1</m:t></m:r></m:num><m:den><m:r><m:t>2</m:t></m:r></m:den></m:f>"))
    assert "<mfrac><mrow><mn>1</mn></mrow><mrow><mn>2</mn></mrow></mfrac>" in result


def test_scripts():
    result, _ = convert(
        equation(
            "<m:sSubSup><m:e><m:r><m:t>𝑁</m:t></m:r></m:e><m:sub><m:r><m:t>𝛾</m:t></m:r></m:sub>"
            "<m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSubSup>"
        )
    )
    assert result.startswith("<math><mrow><msubsup><mrow><mi>𝑁</mi></mrow>")


def test_radical():
    square, _ = convert(equation("<m:rad><m:radPr><m:degHide m:val=\"1\"/></m:radPr><m:deg/><m:e><m:r><m:t>𝑥</m:t></m:r></m:e></m:rad>"))
    assert "<msqrt>" in square
    cube, _ = convert(equation("<m:rad><m:deg><m:r><m:t>3</m:t></m:r></m:deg><m:e><m:r><m:t>𝑥</m:t></m:r></m:e></m:rad>"))
    assert "<mroot>" in cube


def test_brackets_stretch_only_around_tall_content():
    short, _ = convert(equation("<m:d><m:e><m:r><m:t>𝑋</m:t></m:r></m:e></m:d>"))
    assert 'stretchy="false">(' in short
    tall, _ = convert(
        equation(
            "<m:d><m:e><m:f><m:num><m:r><m:t>1</m:t></m:r></m:num>"
            "<m:den><m:r><m:t>2</m:t></m:r></m:den></m:f></m:e></m:d>"
        )
    )
    assert 'stretchy="true">(' in tall


def test_angle_brackets():
    result, _ = convert(
        equation('<m:d><m:dPr><m:begChr m:val="⟨"/><m:endChr m:val="⟩"/></m:dPr><m:e><m:r><m:t>𝑋</m:t></m:r></m:e></m:d>')
    )
    assert ">⟨</mo>" in result
    assert ">⟩</mo>" in result


def test_nary():
    result, _ = convert(
        equation(
            '<m:nary><m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr>'
            "<m:sub><m:r><m:t>𝑖</m:t></m:r></m:sub><m:sup><m:r><m:t>𝑛</m:t></m:r></m:sup>"
            "<m:e><m:r><m:t>𝑥</m:t></m:r></m:e></m:nary>"
        )
    )
    assert "<munderover><mo largeop=\"true\">∑</mo>" in result


def test_display_equation():
    element = etree.fromstring(
        f'<a14:m xmlns:a14="{A14}" xmlns:m="{M}"><m:oMathPara><m:oMath>'
        "<m:r><m:t>𝑥</m:t></m:r></m:oMath></m:oMathPara></a14:m>"
    )
    result, _ = convert(element)
    assert result.startswith('<math display="block">')


def test_unknown_structure_warns():
    result, warnings = convert(equation("<m:sometime><m:r><m:t>𝑥</m:t></m:r></m:sometime>"))
    assert "<mtext>𝑥</mtext>" in result
    assert warnings == ["an equation uses m:sometime, which is shown as plain text"]


def deck_with_box() -> tuple[Deck, Slide, Shape]:
    presentation = Presentation()
    presentation.slide_width = Emu(WIDTH)
    presentation.slide_height = Emu(HEIGHT)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(WIDTH // 2), Emu(HEIGHT // 2))
    box.text_frame.word_wrap = True
    box.text_frame.text = "If "
    return presentation, slide, box


def page(presentation: Deck, folder: pathlib.Path) -> html.HtmlElement:
    path = folder / "deck.pptx"
    presentation.save(str(path))
    return html.parse(str(pptx_to_html.convert(path, folder / "out", title="Talk"))).getroot()


def test_equation_in_a_text_box(tmp_path):
    presentation, _, box = deck_with_box()
    paragraph = cast(etree._Element, box.text_frame.paragraphs[0]._p)
    paragraph.append(equation("<m:r><m:t>𝑋</m:t></m:r>"))

    result = page(presentation, tmp_path)

    text = result.find_class("text")[0]
    assert text.find(".//math") is not None
    assert "If" in text.text_content()


def test_newer_version_is_drawn_instead_of_its_stand_in(tmp_path):
    presentation, slide, box = deck_with_box()
    paragraph = cast(etree._Element, box.text_frame.paragraphs[0]._p)
    paragraph.append(equation("<m:r><m:t>𝑋</m:t></m:r>"))
    _, rid = slide.part.get_or_add_image_part(_png())
    stand_in = etree.fromstring(
        f'<p:sp xmlns:p="{P}" xmlns:a="{A}" xmlns:r="{R}"><p:nvSpPr><p:cNvPr id="99" name="Stand-in"/>'
        "<p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x=\"0\" y=\"0\"/><a:ext cx=\"100\" cy=\"100\"/>"
        f'</a:xfrm><a:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></a:blipFill>'
        "</p:spPr></p:sp>"
    )
    alternative = etree.Element(f"{{{MC}}}AlternateContent", nsmap={"mc": MC, "a14": A14})
    choice = etree.SubElement(alternative, f"{{{MC}}}Choice", Requires="a14")
    fallback = etree.SubElement(alternative, f"{{{MC}}}Fallback")
    shape = cast(etree._Element, box._element)
    tree = shape.getparent()
    assert tree is not None
    tree.replace(shape, alternative)
    choice.append(shape)
    fallback.append(stand_in)

    result = page(presentation, tmp_path)

    assert result.find(".//math") is not None
    assert not result.find_class("crop")


def test_picture_fill(tmp_path):
    presentation, slide, box = deck_with_box()
    _, rid = slide.part.get_or_add_image_part(_png())
    shape = cast(etree._Element, box._element)
    properties = shape.find(f"{{{P}}}spPr")
    assert properties is not None
    for old in properties.findall(f"{{{A}}}noFill"):
        properties.remove(old)
    fill = etree.SubElement(properties, f"{{{A}}}blipFill")
    etree.SubElement(fill, f"{{{A}}}blip").set(f"{{{R}}}embed", rid)
    etree.SubElement(etree.SubElement(fill, f"{{{A}}}stretch"), f"{{{A}}}fillRect", l="-10000", t="0", r="0", b="0")

    result = page(presentation, tmp_path)

    piece = result.find_class("crop")[0]
    image = piece.find("img")
    assert image is not None
    assert "left:-10.0000%" in image.get("style", "")
    assert "width:110.0000%" in image.get("style", "")


def _png() -> io.BytesIO:
    buffer = io.BytesIO()
    Image.new("RGB", (4, 4), (0, 0, 0)).save(buffer, "PNG")
    buffer.seek(0)
    return buffer


@pytest.mark.parametrize("typeface", ["Aptos", "Aptos Display"])
def test_stand_ins(typeface):
    stack = font_stack(typeface)
    assert stack.startswith(f'"{typeface}", "{typeface} (Segoe UI)"')
    rules = font_faces(f"font-family:{font_stack(typeface, quote='&#x27;')}")
    assert f'font-family: "{typeface} (Arial)"' in rules
    assert rules.count("@font-face") == 12
    assert "size-adjust" in rules


def test_twins():
    assert '"Carlito"' in font_stack("Calibri")
    assert font_faces(font_stack("Calibri")) == ""
